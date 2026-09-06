import io
import zipfile

import pytest

import attachment_readers as ar


def _xlsx(rows_by_sheet: dict) -> bytes:
    """Builds a real xlsx rather than mocking openpyxl, so these tests exercise
    the actual parse path. A mock here would pass even if the reader were
    pointed at the wrong library entirely."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in rows_by_sheet.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_content_and_sheet_names_come_through():
    raw = _xlsx({"הדרכות": [["חנות", "סטטוס"], ["רחובות", "בוצע"]]})
    out = ar.extract_text("Z8 Training Status.xlsx", raw)
    assert "הדרכות" in out
    assert "רחובות" in out and "בוצע" in out


def test_xlsx_reads_every_sheet_not_just_the_first():
    raw = _xlsx({"ראשון": [["א"]], "שני": [["ב"]], "שלישי": [["ג"]]})
    out = ar.extract_text("multi.xlsx", raw)
    assert "ב" in out and "ג" in out, "later sheets were dropped"


def test_xlsx_no_longer_stops_at_a_row_cap():
    """The 200-row cap was removed on purpose: it hid data the owner had asked
    for, with no way to reach it. A long sheet is now paged, and every row is
    reachable through some part."""
    rows = 600
    raw = _xlsx({"ארוך": [[f"שורה {i}"] for i in range(rows)]})
    first = ar.extract_text("long.xlsx", raw)
    total = int(first.split("מתוך ")[1].split("]")[0]) if "מתוך " in first else 1
    seen = "".join(ar.extract_text("long.xlsx", raw, part=n) for n in range(1, total + 1))
    # The old cap stopped at 200; every one of the 600 has to be reachable now.
    for i in (0, 199, 200, rows - 1):
        assert f"שורה {i}" in seen, f"row {i} is unreachable"


def test_csv_detects_a_semicolon_separator():
    """Hebrew Excel exports semicolons; assuming a comma collapses every row
    into a single column and silently destroys the table."""
    out = ar.extract_text("t.csv", "חנות;מכירות\nלוד;12\n".encode("utf-8"))
    assert "חנות | מכירות" in out
    assert "לוד | 12" in out


def test_csv_decodes_windows_hebrew_encoding():
    out = ar.extract_text("t.csv", "שלום,עולם".encode("cp1255"))
    assert "שלום" in out, "cp1255 was decoded as something else"


def test_plain_text_is_returned_as_is():
    assert "היי" in ar.extract_text("note.txt", "היי".encode("utf-8"))


def test_docx_includes_text_held_in_tables():
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("פסקה רגילה")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "יבנה"
    table.rows[0].cells[1].text = "הושלם"
    buffer = io.BytesIO()
    document.save(buffer)
    out = ar.extract_text("doc.docx", buffer.getvalue())
    assert "פסקה רגילה" in out
    assert "יבנה" in out and "הושלם" in out


def test_a_scanned_pdf_is_explained_not_reported_as_empty():
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    out = ar.extract_text("scan.pdf", buffer.getvalue())
    assert "OCR" in out or "סרוק" in out


def test_oversized_files_are_refused_before_parsing():
    out = ar.extract_text("huge.xlsx", b"x" * (ar.MAX_ATTACHMENT_BYTES + 1))
    assert out.startswith("❌") and "גדול" in out


def test_one_page_stays_within_the_prompt_budget():
    """Output is still bounded per call - the file is paged, not dumped - but the
    bound is now a page rather than a wall, so the rest stays reachable."""
    out = ar.extract_text("big.txt", ("א" * 50_000).encode("utf-8"))
    # The page itself is capped; the header and footer are small fixed additions.
    assert len(out) <= ar.MAX_TEXT_CHARS + 200
    assert "אפשר לבקש את חלק 2" in out


def test_old_excel_format_gets_a_useful_explanation():
    out = ar.extract_text("old.xls", b"\xd0\xcf\x11\xe0")
    assert "xlsx" in out, "the user is not told how to fix it"


def test_images_are_declined_clearly():
    out = ar.extract_text("photo.jpg", b"\xff\xd8\xff")
    assert out.startswith("❌") and "תמונה" in out


def test_a_corrupt_zip_container_does_not_raise():
    out = ar.extract_text("broken.xlsx", b"not a zip at all")
    assert out.startswith("❌")
    assert "פגום" in out or "לא הצלחתי" in out


def test_unknown_types_list_what_is_supported():
    out = ar.extract_text("thing.dwg", b"abc")
    assert ".xlsx" in out and ".pdf" in out


def test_extension_beats_a_wrong_mime_type():
    """Forwarded spreadsheets routinely arrive as application/octet-stream;
    trusting the MIME type would make them unreadable."""
    raw = _xlsx({"גיליון": [["ערך"]]})
    out = ar.extract_text("s.xlsx", raw, mime_type="application/octet-stream")
    assert "ערך" in out


def test_mime_type_is_used_when_there_is_no_extension():
    out = ar.extract_text("attachment", "טקסט".encode("utf-8"), mime_type="text/plain")
    assert "טקסט" in out


def test_is_supported_answers_from_the_name_alone():
    assert ar.is_supported("a.xlsx") and ar.is_supported("b.pdf")
    assert not ar.is_supported("c.jpg") and not ar.is_supported("d.xls")


def test_empty_attachment_is_reported():
    assert ar.extract_text("e.txt", b"").startswith("❌")


def test_a_valid_zip_that_is_not_an_xlsx_fails_gracefully():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("hello.txt", "hi")
    assert ar.extract_text("fake.xlsx", buffer.getvalue()).startswith("❌")
